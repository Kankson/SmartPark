from __future__ import annotations

import html
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_blocks(document: Document):
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def paragraph_html(paragraph: Paragraph) -> str:
    text = html.escape(paragraph.text)
    style = paragraph.style.name if paragraph.style else "Normal"
    page_break = bool(paragraph._p.xpath(".//w:br[@w:type='page']"))
    break_html = '<div class="page-break"></div>' if page_break else ""
    if not text:
        return break_html
    if style.startswith("Heading 1"):
        return break_html + f"<h1>{text}</h1>"
    if style.startswith("Heading 2"):
        return break_html + f"<h2>{text}</h2>"
    if style.startswith("Heading 3"):
        return break_html + f"<h3>{text}</h3>"
    if style.startswith("List Number"):
        return break_html + f'<p class="list number">{text}</p>'
    if style.startswith("List Bullet"):
        return break_html + f'<p class="list bullet">{text}</p>'
    if any(run.font.name == "Consolas" for run in paragraph.runs):
        return break_html + f"<pre>{text}</pre>"
    return break_html + f"<p>{text}</p>"


def table_html(table: Table) -> str:
    rows = []
    for row_index, row in enumerate(table.rows):
        tag = "th" if row_index == 0 else "td"
        cells = "".join(f"<{tag}>{html.escape(cell.text)}</{tag}>" for cell in row.cells)
        rows.append(f"<tr>{cells}</tr>")
    return "<table>" + "".join(rows) + "</table>"


def main() -> None:
    source = Path(sys.argv[1]).resolve()
    target = Path(sys.argv[2]).resolve()
    document = Document(source)
    content = []
    for block in iter_blocks(document):
        content.append(paragraph_html(block) if isinstance(block, Paragraph) else table_html(block))

    markup = f"""<!doctype html>
<html><head><meta charset="utf-8"><title>SmartPark Documentation QA</title>
<style>
@page {{ size: Letter; margin: 1in; }}
* {{ box-sizing: border-box; }}
body {{ margin: 0; color: #111827; font: 11pt Calibri, Arial, sans-serif; line-height: 1.45; }}
h1 {{ color: #173f5f; font-size: 19pt; margin: 16pt 0 8pt; break-after: avoid; }}
h2 {{ color: #173f5f; font-size: 15pt; margin: 12pt 0 6pt; break-after: avoid; }}
h3 {{ color: #173f5f; font-size: 12pt; margin: 9pt 0 4pt; break-after: avoid; }}
p {{ margin: 0 0 7pt; text-align: justify; }}
.list {{ padding-left: 20pt; position: relative; text-align: left; }}
.number::before {{ content: counter(item) '. '; counter-increment: item; position: absolute; left: 0; }}
.bullet::before {{ content: '\\2022'; position: absolute; left: 2pt; }}
pre {{ margin: 0 0 3pt 18pt; color: #4b5563; font: 9pt Consolas, monospace; white-space: pre-wrap; }}
table {{ border-collapse: collapse; width: 100%; margin: 5pt 0 10pt; font-size: 9pt; break-inside: auto; }}
tr {{ break-inside: avoid; }}
th, td {{ border: 1px solid #6b7280; padding: 4pt; vertical-align: top; }}
th {{ background: #e5e7eb; text-align: left; }}
.page-break {{ break-after: page; }}
@media screen {{ body {{ max-width: 8.5in; margin: 0 auto; padding: 1in; background: white; }} }}
</style></head><body>{''.join(content)}</body></html>"""
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(markup, encoding="utf-8")
    print(target)
    print(f"paragraphs={len(document.paragraphs)} tables={len(document.tables)}")


if __name__ == "__main__":
    main()
