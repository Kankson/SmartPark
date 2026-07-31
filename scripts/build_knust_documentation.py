from __future__ import annotations

import argparse
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "knust-condensed-report.md"
ASSET_DIR = ROOT / "DOCUMENTATION" / "assets" / "knust_diagrams"
DEFAULT_OUT = ROOT / "DOCUMENTATION" / "SmartPark_KNUST_Project_Report_Under_30_Pages.docx"

TNR = "Times New Roman"
BLACK = RGBColor(0, 0, 0)
GREEN = RGBColor(0, 107, 63)
PALE_GREEN = "E2F0E9"
BORDER = "7F7F7F"
USABLE_DXA = 8220

FIGURES = [
    ("Figure 2.1: Conceptual framework for SmartPark", "figure_2_1_conceptual_framework.png"),
    ("Figure 3.1: SmartPark system architecture", "figure_3_1_system_architecture.png"),
    ("Figure 3.2: SmartPark use-case model", "figure_3_2_use_case_model.png"),
    ("Figure 3.3: Core SmartPark data model", "figure_3_3_core_data_model.png"),
    ("Figure 4.1: End-to-end SmartPark parking workflow", "figure_4_1_end_to_end_workflow.png"),
    ("Figure 4.2: Booking and enforcement lifecycle", "figure_4_2_booking_lifecycle.png"),
]

TABLES = [
    "Table 2.1: Comparison of selected parking approaches",
    "Table 3.1: Consolidated system requirements",
    "Table 4.1: Main development tools",
    "Table 4.2: Verification results",
]

TOC_ENTRIES = [
    (0, "DECLARATION", "ii"),
    (0, "ABSTRACT", "iii"),
    (0, "ACKNOWLEDGEMENTS", "iv"),
    (0, "TABLE OF CONTENTS", "v"),
    (0, "LIST OF TABLES", "vi"),
    (0, "LIST OF FIGURES", "vii"),
    (0, "LIST OF ABBREVIATIONS", "viii"),
    (0, "CHAPTER ONE: INTRODUCTION", "CHAPTER ONE"),
    (1, "1.1 Background to the Study", "1.1 Background to the Study"),
    (1, "1.2 Problem Statement", "1.2 Problem Statement"),
    (1, "1.3 Aim of the Project", "1.3 Aim of the Project"),
    (1, "1.4 Project Objectives", "1.4 Project Objectives"),
    (1, "1.5 Research Questions", "1.5 Research Questions"),
    (1, "1.6 Significance of the Project", "1.6 Significance of the Project"),
    (1, "1.7 Scope and Delimitations", "1.7 Scope and Delimitations"),
    (1, "1.8 Organisation of the Report", "1.8 Organisation of the Report"),
    (0, "CHAPTER TWO: REVIEW OF RELATED LITERATURE", "CHAPTER TWO"),
    (1, "2.1 Introduction", "2.1 Introduction"),
    (1, "2.2 Digital Parking and Mobile Payment", "2.2 Digital Parking and Mobile Payment"),
    (1, "2.3 QR-Based Parking Access", "2.3 QR-Based Parking Access"),
    (1, "2.4 Availability Without Space Sensors", "2.4 Availability Without Space Sensors"),
    (1, "2.5 Mobile Enforcement and Human Oversight", "2.5 Mobile Enforcement and Human Oversight"),
    (1, "2.6 Existing Systems and Identified Gap", "2.6 Existing Systems and Identified Gap"),
    (1, "2.7 Chapter Summary", "2.7 Chapter Summary"),
    (0, "CHAPTER THREE: METHODOLOGY, ANALYSIS AND DESIGN", "CHAPTER THREE"),
    (1, "3.1 Introduction", "3.1 Introduction"),
    (1, "3.2 Development Methodology", "3.2 Development Methodology"),
    (1, "3.3 Functional and Non-Functional Requirements", "3.3 Functional and Non-Functional Requirements"),
    (1, "3.4 System Architecture", "3.4 System Architecture"),
    (1, "3.5 Use-Case Model", "3.5 Use-Case Model"),
    (1, "3.6 Data Design", "3.6 Data Design"),
    (1, "3.7 Security, Privacy and Ethical Considerations", "3.7 Security, Privacy and Ethical Considerations"),
    (1, "3.8 Chapter Summary", "3.8 Chapter Summary"),
    (0, "CHAPTER FOUR: PROJECT EXECUTION, IMPLEMENTATION AND TESTING", "CHAPTER FOUR"),
    (1, "4.1 Introduction", "4.1 Introduction"),
    (1, "4.2 Development Environment", "4.2 Development Environment"),
    (1, "4.3 Implemented Workflow", "4.3 Implemented Workflow"),
    (1, "4.4 Booking and Enforcement Lifecycle", "4.4 Booking and Enforcement Lifecycle"),
    (1, "4.5 Driver, Warden and Administrator Modules", "4.5 Driver, Warden and Administrator Modules"),
    (1, "4.6 Verification and Test Results", "4.6 Verification and Test Results"),
    (1, "4.7 Results and Discussion", "4.7 Results and Discussion"),
    (1, "4.8 Deployment Considerations", "4.8 Deployment Considerations"),
    (1, "4.9 Chapter Summary", "4.9 Chapter Summary"),
    (0, "CHAPTER FIVE: SUMMARY, CONCLUSIONS AND RECOMMENDATIONS", "CHAPTER FIVE"),
    (1, "5.1 Summary", "5.1 Summary"),
    (1, "5.2 Conclusions", "5.2 Conclusions"),
    (1, "5.3 Recommendations", "5.3 Recommendations"),
    (1, "5.4 Future Work", "5.4 Future Work"),
    (0, "REFERENCES", "REFERENCES"),
    (0, "APPENDIX A: DEMONSTRATION AND REPRODUCTION DETAILS", "APPENDIX A"),
]


def font_run(run, size: float = 12, bold: bool = False, italic: bool = False) -> None:
    run.font.name = TNR
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), TNR)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), TNR)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = BLACK


def configure_section(section) -> None:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(4.0)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = TNR
    normal._element.rPr.rFonts.set(qn("w:ascii"), TNR)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), TNR)
    normal.font.size = Pt(12)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = True

    heading1 = styles["Heading 1"]
    heading1.font.name = TNR
    heading1._element.rPr.rFonts.set(qn("w:ascii"), TNR)
    heading1._element.rPr.rFonts.set(qn("w:hAnsi"), TNR)
    heading1.font.size = Pt(14)
    heading1.font.bold = True
    heading1.font.color.rgb = BLACK
    heading1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading1.paragraph_format.line_spacing = 1.0
    heading1.paragraph_format.space_before = Pt(0)
    heading1.paragraph_format.space_after = Pt(14)
    heading1.paragraph_format.keep_with_next = True

    heading2 = styles["Heading 2"]
    heading2.font.name = TNR
    heading2._element.rPr.rFonts.set(qn("w:ascii"), TNR)
    heading2._element.rPr.rFonts.set(qn("w:hAnsi"), TNR)
    heading2.font.size = Pt(12)
    heading2.font.bold = True
    heading2.font.color.rgb = BLACK
    heading2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    heading2.paragraph_format.line_spacing = 1.0
    heading2.paragraph_format.space_before = Pt(12)
    heading2.paragraph_format.space_after = Pt(6)
    heading2.paragraph_format.keep_with_next = True

    for name in ["Caption"]:
        style = styles[name]
        style.font.name = TNR
        style._element.rPr.rFonts.set(qn("w:ascii"), TNR)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), TNR)
        style.font.size = Pt(10)
        style.font.color.rgb = BLACK
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.space_before = Pt(3)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = False


def set_page_number_format(section, fmt: str, start: int) -> None:
    sect_pr = section._sectPr
    existing = sect_pr.find(qn("w:pgNumType"))
    if existing is not None:
        sect_pr.remove(existing)
    pg_num = OxmlElement("w:pgNumType")
    pg_num.set(qn("w:fmt"), fmt)
    pg_num.set(qn("w:start"), str(start))
    sect_pr.append(pg_num)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    font_run(run, size=10)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def setup_footer(section, fmt: str, start: int) -> None:
    section.footer.is_linked_to_previous = False
    footer = section.footer
    p = footer.paragraphs[0]
    p.clear()
    add_page_field(p)
    set_page_number_format(section, fmt, start)


def add_inline_runs(paragraph, text: str, size: float = 12) -> None:
    cursor = 0
    pattern = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`.+?`)")
    for match in pattern.finditer(text):
        if match.start() > cursor:
            font_run(paragraph.add_run(text[cursor : match.start()]), size=size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            font_run(run, size=size, bold=True)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            font_run(run, size=size, italic=True)
        else:
            run = paragraph.add_run(token[1:-1])
            font_run(run, size=size)
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        cursor = match.end()
    if cursor < len(text):
        font_run(paragraph.add_run(text[cursor:]), size=size)


def add_caption_runs(paragraph, text: str) -> None:
    label, separator, description = text.partition(":")
    font_run(paragraph.add_run(label + separator), size=10, bold=True)
    if description:
        font_run(paragraph.add_run(description), size=10)


def add_body_paragraph(doc: Document, text: str, *, indent: bool = True, single: bool = False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.0 if single else 1.5
    p.paragraph_format.space_after = Pt(4 if single else 0)
    if indent:
        p.paragraph_format.first_line_indent = Cm(1.0)
    add_inline_runs(p, text)
    return p


def numbering_id(doc: Document, fmt: str, marker: str) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), marker)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "space")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.extend([tabs, ind])
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), TNR)
    r_fonts.set(qn("w:hAnsi"), TNR)
    r_pr.append(r_fonts)
    lvl.extend([start, num_fmt, lvl_text, suff, p_pr, r_pr])
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_list_paragraph(doc: Document, text: str, num_id: int) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(2)
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])
    p_pr.append(num_pr)
    add_inline_runs(p, text)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="5") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths: list[int], indent: int = 120) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
    if tbl_w.getparent() is None:
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
    if tbl_ind.getparent() is None:
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[min(index, len(widths) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
            if tc_w.getparent() is None:
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    cols = max(len(row) for row in rows)
    if cols == 2:
        widths = [2300, USABLE_DXA - 2300]
    elif cols == 3:
        widths = [1900, 2600, USABLE_DXA - 4500]
    else:
        widths = [USABLE_DXA // cols] * cols
        widths[-1] = USABLE_DXA - sum(widths[:-1])
    table = doc.add_table(rows=len(rows), cols=cols)
    set_table_geometry(table, widths)
    set_table_borders(table)
    header_props = table.rows[0]._tr.get_or_add_trPr()
    header_tag = OxmlElement("w:tblHeader")
    header_tag.set(qn("w:val"), "true")
    header_props.append(header_tag)
    for row_index, row in enumerate(rows):
        for col_index in range(cols):
            cell = table.cell(row_index, col_index)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            add_inline_runs(p, row[col_index] if col_index < len(row) else "", size=9.5)
            if row_index == 0:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), PALE_GREEN)
                cell._tc.get_or_add_tcPr().append(shd)
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = GREEN
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_borderless_index_table(doc: Document, rows: list[tuple[str, str]], font_size: float = 10, indent_levels: list[int] | None = None) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_geometry(table, [USABLE_DXA - 700, 700], indent=0)
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    tbl_pr.append(borders)
    for index, (label, page) in enumerate(rows):
        left, right = table.rows[index].cells
        for cell in (left, right):
            cell.text = ""
            set_cell_margins(cell, 0, 0, 0, 0)
        p = left.paragraphs[0]
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        if indent_levels and indent_levels[index]:
            p.paragraph_format.left_indent = Cm(0.5)
        add_inline_runs(p, label, size=font_size)
        if not indent_levels or not indent_levels[index]:
            for run in p.runs:
                run.bold = True
        p = right.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        add_inline_runs(p, str(page), size=font_size)


def add_center_heading(doc: Document, text: str, size: float = 14) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(12)
    font_run(p.add_run(text), size=size, bold=True)


def add_title_page(doc: Document) -> None:
    def cover_line(
        text: str,
        *,
        size: float = 12,
        before: float = 0,
        after: float = 8,
        color: RGBColor = BLACK,
    ) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        run = p.add_run(text)
        font_run(run, size=size, bold=True)
        run.font.color.rgb = color

    cover_line("KWAME NKRUMAH UNIVERSITY OF SCIENCE AND TECHNOLOGY", size=14, before=6, after=12)
    cover_line("COLLEGE OF SCIENCE", size=13, after=10)
    cover_line("FACULTY OF PHYSICAL AND COMPUTATIONAL SCIENCE", size=13, after=18)
    cover_line("PROJECT DOCUMENTATION", size=15, after=12, color=GREEN)
    cover_line("FOR CONSIDERATION UNDER", size=12, after=10)
    cover_line("BACHELOR OF COMPUTER SCIENCE", size=13, after=4)
    cover_line("(BSc) PROGRAMME", size=12, after=16)
    cover_line("PROJECT TITLE", size=12, after=8)
    cover_line("SMARTPARK: A SIMPLIFIED SMART PARKING MANAGEMENT SYSTEM", size=16, after=16, color=GREEN)
    cover_line("SUBMITTED BY", size=12, after=8)

    for name, index_no, student_id in [
        ("KPENTEY KANKSON KOMLA EWOENAM", "3396222", "20940297"),
        ("WILLIAM JOHNSON", "3394222", "20938022"),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(7)
        font_run(p.add_run(name + "\n"), size=12, bold=True)
        font_run(p.add_run(f"INDEX NUMBER: {index_no}  |  STUDENT ID: {student_id}"), size=10.5)

    cover_line("DR. KWABENA AGYEMANG", size=12, before=10, after=3)
    cover_line("(Supervisor)", size=11, after=0)


def add_declaration(doc: Document) -> None:
    add_center_heading(doc, "DECLARATION")
    statement = (
        "We hereby declare that this submission is our own work and that, to the best of our knowledge "
        "and belief, it contains no material previously published or written by another person, nor "
        "material which to a substantial extent has been accepted for the award of any other degree or "
        "diploma at Kwame Nkrumah University of Science and Technology, Kumasi, or any other educational "
        "institution, except where due acknowledgement is made in the report."
    )
    add_body_paragraph(doc, statement, indent=False)
    doc.add_paragraph()
    lines = [
        ("Kpentey Kankson Komla Ewoenam\nIndex No. 3396222 | Student ID 20940297", "Signature", "Date"),
        ("William Johnson\nIndex No. 3394222 | Student ID 20938022", "Signature", "Date"),
        ("Certified by: Dr. Kwabena Agyemang", "Signature", "Date"),
        ("Certified by: Head of Department", "Signature", "Date"),
    ]
    table = doc.add_table(rows=len(lines), cols=3)
    set_table_geometry(table, [4300, 1900, 1900], indent=0)
    for row, values in zip(table.rows, lines):
        for cell, value in zip(row.cells, values):
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            font_run(p.add_run("________________________\n" + value), size=9.5)


def add_abstract(doc: Document) -> None:
    add_center_heading(doc, "ABSTRACT")
    text = (
        "SmartPark is a phone-first parking management system developed to improve city-centre fee "
        "collection and verification without requiring IoT sensors, automated barriers or fixed cameras. "
        "The project used iterative prototyping and a layered software architecture. Drivers can scan a "
        "signed parking-zone QR code or use a simplified map, compare availability confidence, reserve a "
        "space, pay through a mock wallet, receive a secure ticket QR code, monitor the official end time "
        "and extend an eligible session. Wardens can validate entry and exit, search or capture number "
        "plates, review recognition confidence and resolve priority-ranked violations. Administrators can "
        "view setup data and print signed zone signs. The implementation uses Next.js, React, TypeScript, "
        "MapLibre and a Supabase-ready data model. Verification comprised TypeScript checking, ESLint, "
        "nine automated tests, a production build and selected mobile-browser journeys at a 390 by 844 "
        "pixel viewport; all checks passed. The study concludes that an auditable booking-to-enforcement "
        "workflow can be demonstrated with ordinary phones and printed signs. Production use still "
        "requires persistent storage, the supplied regulated payment provider, validated plate recognition, "
        "field testing and a formal privacy and security review."
    )
    p = add_body_paragraph(doc, text, indent=False, single=True)
    p.paragraph_format.space_after = Pt(10)
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    font_run(p.add_run("Keywords: "), size=11, bold=True)
    font_run(p.add_run("smart parking; digital payment; QR ticketing; parking enforcement; progressive web application"), size=11, italic=True)


def add_acknowledgements(doc: Document) -> None:
    add_center_heading(doc, "ACKNOWLEDGEMENTS")
    text = (
        "We are grateful to Almighty God for the strength and guidance to complete this project. We thank "
        "our supervisor, Dr. Kwabena Agyemang, for his direction, constructive criticism and encouragement. "
        "We also appreciate the lecturers of the Department of Computer Science, our colleagues, friends and "
        "families for the technical suggestions and support provided during the design, implementation and "
        "documentation of SmartPark."
    )
    add_body_paragraph(doc, text, indent=False)


def page_value(page_map: dict[str, str | int], key: str, fallback: str = "...") -> str:
    return str(page_map.get(key, fallback))


def add_toc(doc: Document, page_map: dict[str, str | int]) -> None:
    add_center_heading(doc, "TABLE OF CONTENTS")
    rows = []
    levels = []
    for level, label, key in TOC_ENTRIES:
        page = key if re.fullmatch(r"[ivx]+", key) else page_value(page_map, key)
        rows.append((label, page))
        levels.append(level)
    add_borderless_index_table(doc, rows, font_size=9.2, indent_levels=levels)


def add_list_of_tables(doc: Document, page_map: dict[str, str | int]) -> None:
    add_center_heading(doc, "LIST OF TABLES")
    add_borderless_index_table(doc, [(caption, page_value(page_map, caption)) for caption in TABLES], font_size=10.5)


def add_list_of_figures(doc: Document, page_map: dict[str, str | int]) -> None:
    add_center_heading(doc, "LIST OF FIGURES")
    add_borderless_index_table(doc, [(caption, page_value(page_map, caption)) for caption, _ in FIGURES], font_size=10.5)


def add_abbreviations(doc: Document) -> None:
    add_center_heading(doc, "LIST OF ABBREVIATIONS")
    rows = [
        ("ANPR", "Automatic Number Plate Recognition"),
        ("API", "Application Programming Interface"),
        ("HMAC", "Hash-Based Message Authentication Code"),
        ("IoT", "Internet of Things"),
        ("MVP", "Minimum Viable Product"),
        ("OCR", "Optical Character Recognition"),
        ("PWA", "Progressive Web Application"),
        ("QR", "Quick Response"),
        ("RLS", "Row-Level Security"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_geometry(table, [1600, USABLE_DXA - 1600], indent=0)
    for row, values in zip(table.rows, rows):
        for index, (cell, value) in enumerate(zip(row.cells, values)):
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.5
            add_inline_runs(p, value, size=11)
            if index == 0:
                for run in p.runs:
                    run.bold = True


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    raw_rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        raw_rows.append(lines[index].strip())
        index += 1
    rows = []
    for line in raw_rows:
        cells = [cell.strip() for cell in line.strip("|").split("|")]
        if all(re.fullmatch(r"-+", cell.replace(" ", "")) for cell in cells):
            continue
        rows.append(cells)
    return rows, index


def add_figure(doc: Document, filename: str, caption: str) -> None:
    path = ASSET_DIR / filename
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=Inches(5.55))
    drawing = run._r.xpath(".//wp:docPr")
    if drawing:
        drawing[0].set("descr", caption)
        drawing[0].set("title", caption.split(":", 1)[0])
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    add_caption_runs(cap, caption)


def add_main_content(doc: Document) -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    first_top_level = True
    pending_table_caption: str | None = None
    active_num_id: int | None = None

    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped:
            active_num_id = None
            index += 1
            continue

        table_marker = re.fullmatch(r"\[\[TABLE:(.+)\]\]", stripped)
        if table_marker:
            pending_table_caption = table_marker.group(1).strip()
            index += 1
            continue

        figure_marker = re.fullmatch(r"\[\[FIGURE:([^|]+)\|(.+)\]\]", stripped)
        if figure_marker:
            add_figure(doc, figure_marker.group(1).strip(), figure_marker.group(2).strip())
            index += 1
            continue

        if stripped.startswith("|"):
            rows, index = parse_table(lines, index)
            if pending_table_caption:
                cap = doc.add_paragraph(style="Caption")
                cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
                cap.paragraph_format.keep_with_next = True
                add_caption_runs(cap, pending_table_caption)
                pending_table_caption = None
            add_table(doc, rows)
            continue

        if stripped.startswith("# "):
            title = stripped[2:].strip()
            if title.startswith("CHAPTER ") or title in {"REFERENCES", "APPENDIX A"}:
                starts_new_page = not first_top_level
                first_top_level = False
                p = doc.add_paragraph(style="Heading 1")
                p.paragraph_format.page_break_before = starts_new_page
                p.paragraph_format.space_after = Pt(6)
                add_inline_runs(p, title, size=14)
            else:
                p = doc.add_paragraph(style="Heading 1")
                add_inline_runs(p, title, size=14)
            index += 1
            continue

        if stripped.startswith("## "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline_runs(p, stripped[3:].strip(), size=12)
            index += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if numbered:
            if active_num_id is None:
                active_num_id = numbering_id(doc, "decimal", "%1.")
            add_list_paragraph(doc, numbered.group(1), active_num_id)
            index += 1
            continue

        active_num_id = None
        add_body_paragraph(doc, stripped, indent=not stripped.startswith("http"), single=stripped.startswith("http"))
        index += 1


def scrub_metadata(path: Path) -> None:
    tmp = path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/core.xml":
                text = data.decode("utf-8")
                text = re.sub(r"<dc:creator>.*?</dc:creator>", "<dc:creator>SmartPark Project Team</dc:creator>", text)
                text = re.sub(r"<cp:lastModifiedBy>.*?</cp:lastModifiedBy>", "<cp:lastModifiedBy>SmartPark Project Team</cp:lastModifiedBy>", text)
                data = text.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)


def build(output: Path, page_map: dict[str, str | int]) -> None:
    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0])
    doc.sections[0].footer.is_linked_to_previous = False
    doc.sections[0].footer.paragraphs[0].clear()
    add_title_page(doc)

    front = doc.add_section(WD_SECTION_START.NEW_PAGE)
    configure_section(front)
    setup_footer(front, "lowerRoman", 2)
    add_declaration(doc)
    doc.add_page_break()
    add_abstract(doc)
    doc.add_page_break()
    add_acknowledgements(doc)
    doc.add_page_break()
    add_toc(doc, page_map)
    doc.add_page_break()
    add_list_of_tables(doc, page_map)
    doc.add_page_break()
    add_list_of_figures(doc, page_map)
    doc.add_page_break()
    add_abbreviations(doc)

    main = doc.add_section(WD_SECTION_START.NEW_PAGE)
    configure_section(main)
    setup_footer(main, "decimal", 1)
    add_main_content(doc)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    scrub_metadata(output)
    print(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    parser.add_argument("--page-map", type=Path)
    args = parser.parse_args()
    page_map: dict[str, str | int] = {}
    if args.page_map and args.page_map.exists():
        page_map = json.loads(args.page_map.read_text(encoding="utf-8-sig"))
    build(args.out.resolve(), page_map)


if __name__ == "__main__":
    main()
