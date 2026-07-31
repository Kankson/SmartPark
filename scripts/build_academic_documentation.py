from __future__ import annotations

import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "DOCUMENTATION"
OUT_DOCX = OUT_DIR / "SmartPark_Full_Project_Documentation_Updated.docx"
CHAPTERS = [
    ROOT / "docs" / "chapter-one.md",
    ROOT / "docs" / "chapter-two.md",
    ROOT / "docs" / "chapter-three.md",
    ROOT / "docs" / "chapter-four.md",
    ROOT / "docs" / "chapter-five.md",
]

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
GRAY = RGBColor(88, 88, 88)
LIGHT_GRAY = "F2F4F7"
BORDER = "D9E2EC"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def set_borders(table, color=BORDER, size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_text)
    run._r.append(fld_end)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ["List Bullet", "List Number"]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)


def paragraph(doc: Document, text="", style=None, align=None, bold=False, italic=False, color=None, size=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    if size is not None:
        run.font.size = Pt(size)
    return p


def add_inline_runs(p, text: str, size: int | None = None) -> None:
    cursor = 0
    for match in re.finditer(r"(\*\*.+?\*\*|`.+?`)", text):
        if match.start() > cursor:
            run = p.add_run(text[cursor : match.start()])
            if size is not None:
                run.font.size = Pt(size)
        token = match.group(0)
        run = p.add_run(token[2:-2] if token.startswith("**") else token[1:-1])
        if token.startswith("**"):
            run.bold = True
        else:
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        if size is not None:
            run.font.size = Pt(size)
        cursor = match.end()
    if cursor < len(text):
        run = p.add_run(text[cursor:])
        if size is not None:
            run.font.size = Pt(size)


def add_cover(doc: Document) -> None:
    paragraph(doc, "SMARTPARK", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, color=BLUE, size=24)
    p = paragraph(
        doc,
        "Simplified Smart Parking Management System",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        bold=True,
        size=18,
    )
    p.paragraph_format.space_after = Pt(20)
    paragraph(
        doc,
        "A Project Documentation Submitted in Partial Fulfilment of the Requirements for the Award of [Programme/Degree Name]",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        size=12,
    )
    paragraph(doc, "By", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12)
    paragraph(doc, "[Student Name]", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13)
    paragraph(doc, "[Index Number / Student ID]", align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
    paragraph(doc, "[Department Name]", align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
    paragraph(doc, "[Institution Name]", align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
    paragraph(doc, "Supervisor: [Supervisor Name]", align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
    p = paragraph(doc, "July 2026", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12)
    p.paragraph_format.space_before = Pt(40)
    doc.add_page_break()


def add_front_matter(doc: Document) -> None:
    front_sections = [
        (
            "Declaration",
            [
                "I hereby declare that this project work titled SmartPark: Simplified Smart Parking Management System is my original work and has not been submitted to any institution for the award of a certificate, diploma, or degree.",
                "All sources of information used in this project have been duly acknowledged.",
                "Student Name: ______________________________",
                "Signature: _________________________________",
                "Date: ______________________________________",
                "Supervisor Name: ___________________________",
                "Signature: _________________________________",
                "Date: ______________________________________",
            ],
        ),
        (
            "Acknowledgement",
            [
                "I express my sincere gratitude to Almighty God for the strength, wisdom, and guidance throughout the development of this project.",
                "I also appreciate my supervisor, lecturers, colleagues, friends, and family for their support, encouragement, and useful suggestions during the project work.",
            ],
        ),
        (
            "Dedication",
            [
                "This project is dedicated to my family, mentors, and everyone who supported my academic journey and encouraged the successful completion of this work.",
            ],
        ),
        (
            "Abstract",
            [
                "SmartPark is a simplified, phone-first parking management system for digital fee collection, booking, and verification in a city-centre environment. Drivers can scan a signed parking-zone QR sign or use a map, compare ranked zones with availability-confidence labels, book and pay through a mock wallet, receive a secure ticket QR, monitor the timer, extend an eligible session, and enable browser reminders.",
                "Wardens can monitor space and session status, scan ticket QR codes, perform manual or camera-assisted plate checks, review recognition confidence before verification, and process live violations ranked by overstay severity. Administrators can view system setup and produce printable signed QR signs for parking zones. Human confirmation is required before a recognised plate or violation decision is accepted.",
                "The project is implemented as a responsive Progressive Web Application using Next.js, React, TypeScript, Tailwind CSS, MapLibre, GSAP, and a Supabase-ready data design. It deliberately avoids IoT sensors, automatic barriers, and fixed cameras, relying on ordinary phones, printed signs, booking records, and warden observations. The latest verification passed TypeScript checking, ESLint, nine automated tests, the production build, and selected mobile-browser journeys. Mock payment, mock plate recognition, in-memory storage, and active-browser notifications remain explicit MVP limitations.",
            ],
        ),
    ]
    for title, paras in front_sections:
        paragraph(doc, title, style="Heading 1")
        for text in paras:
            paragraph(doc, text)
        doc.add_page_break()


def add_manual_toc(doc: Document) -> None:
    paragraph(doc, "Table of Contents", style="Heading 1")
    entries = [
        "Declaration",
        "Acknowledgement",
        "Dedication",
        "Abstract",
        "Chapter One: Introduction",
        "  Problem Statement",
        "  Project Aim",
        "  Project Objectives",
        "  Scope of the Project",
        "  Limitations of the Project",
        "  Approach to Project",
        "Chapter Two: Review of Literature and Tools",
        "  Background Review",
        "  Review of Existing Applications",
        "  Problem Identification",
        "  Project Evaluation",
        "  Review of Project-Related Methodologies",
        "Chapter Three: Requirements Specifications",
        "  Requirements Gathering",
        "  Functional Requirements",
        "  Non-Functional Requirements",
        "  Hardware Requirements",
        "  Requirements Analysis",
        "  Use Case Diagram",
        "Chapter Four: Design Specifications",
        "  System Design and Methodology",
        "  System Interfaces",
        "  Requirement Model",
        "Chapter Five: Implementation and Testing",
        "  Development Tools and Platform Consideration",
        "  System Implementation",
        "  Testing",
        "  Unit Testing",
        "  Integration Testing",
        "  End-to-End Testing",
        "  Deployment",
        "References",
        "Appendices",
    ]
    for entry in entries:
        p = paragraph(doc, entry)
        if entry.startswith("  "):
            p.paragraph_format.left_indent = Inches(0.25)
    doc.add_page_break()


def add_heading(doc: Document, level: int, text: str) -> None:
    if text.strip().lower() == "references":
        paragraph(doc, "References", style="Heading 1")
        return
    paragraph(doc, text.strip(), style=f"Heading {min(level, 3)}")


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    col_count = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=col_count)
    set_borders(table)
    total = 9360
    if col_count == 2:
        widths = [2600, total - 2600]
    elif col_count == 3:
        widths = [2200, 3000, total - 5200]
    elif col_count == 4:
        widths = [1800, 2300, 2700, total - 6800]
    else:
        widths = [int(total / col_count)] * col_count
        widths[-1] = total - sum(widths[:-1])
    set_table_geometry(table, widths)
    header_props = table.rows[0]._tr.get_or_add_trPr()
    header_tag = OxmlElement("w:tblHeader")
    header_tag.set(qn("w:val"), "true")
    header_props.append(header_tag)
    for row_idx, row_data in enumerate(rows):
        for col_idx in range(col_count):
            text = row_data[col_idx] if col_idx < len(row_data) else ""
            cell = table.cell(row_idx, col_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline_runs(p, text, size=10)
            if row_idx == 0:
                for run in p.runs:
                    run.bold = True
                set_cell_shading(cell, LIGHT_GRAY)
    doc.add_paragraph()


def parse_table(lines: list[str], start_index: int):
    table_lines = []
    idx = start_index
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        table_lines.append(lines[idx].strip())
        idx += 1
    rows = []
    for line in table_lines:
        cells = [cell.strip() for cell in line.strip("|").split("|")]
        if all(re.fullmatch(r"-+", cell.replace(" ", "")) for cell in cells):
            continue
        rows.append(cells)
    return rows, idx


def add_code_block(doc: Document, code_lines: list[str]) -> None:
    for line in code_lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        run.font.size = Pt(9)
        run.font.color.rgb = GRAY


def add_markdown_file(doc: Document, path: Path) -> None:
    lines = path.read_text(encoding="utf-8").splitlines()
    i = 0
    in_code = False
    code_lines = []
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not stripped:
            i += 1
            continue
        if stripped.startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            add_heading(doc, level, text)
        elif re.match(r"^\d+\.\s+", stripped):
            text = re.sub(r"^\d+\.\s+", "", stripped)
            p = doc.add_paragraph(style="List Number")
            add_inline_runs(p, text)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, stripped[2:].strip())
        else:
            p = doc.add_paragraph()
            add_inline_runs(p, stripped)
        i += 1

    if code_lines:
        add_code_block(doc, code_lines)


def add_appendices(doc: Document) -> None:
    doc.add_page_break()
    paragraph(doc, "Appendices", style="Heading 1")
    paragraph(doc, "Appendix A: Demo Login Accounts", style="Heading 2")
    add_table(
        doc,
        [
            ["Role", "Email", "Password"],
            ["Driver", "driver@smartpark.test", "password123"],
            ["Driver with active session", "kofi@smartpark.test", "password123"],
            ["Warden", "warden@smartpark.test", "password123"],
            ["Admin", "admin@smartpark.test", "password123"],
        ],
    )
    paragraph(doc, "Appendix B: Local Development Commands", style="Heading 2")
    add_code_block(
        doc,
        [
            'cd "C:\\Users\\Forge Mages\\Documents\\SmartPark"',
            '$env:COREPACK_HOME="$PWD\\.corepack"',
            '$env:CI="true"',
            "corepack pnpm install",
            "corepack pnpm dev -H 127.0.0.1",
            "start http://127.0.0.1:3000",
        ],
    )
    paragraph(doc, "Appendix C: Main Project Files", style="Heading 2")
    add_table(
        doc,
        [
            ["Path", "Purpose"],
            ["src/app", "Application pages and API route handlers"],
            ["src/components", "Reusable user-interface components"],
            ["src/server", "Business logic, domain models, seed data, and providers"],
            ["src/lib", "Authentication, environment, and API helpers"],
            ["supabase/migrations", "Database schema and row-level security"],
            ["docs", "Project documentation chapters"],
        ],
    )
    paragraph(doc, "Appendix D: Latest Verification Results", style="Heading 2")
    add_table(
        doc,
        [
            ["Verification", "Result"],
            ["TypeScript type checking", "Passed"],
            ["ESLint static analysis", "Passed"],
            ["Vitest automated suite", "9 of 9 passed"],
            ["Next.js production build", "Passed"],
            ["Mobile browser journeys", "Passed at 390 x 844"],
        ],
    )
    paragraph(doc, "Appendix E: MVP Boundaries and Production Work", style="Heading 2")
    for item in [
        "The current wallet and payment provider are demonstrations and must be replaced by the supplied regulated payment service before real-money use.",
        "The current data store is in memory; Supabase persistence, authentication, row-level security, backups, and audit retention must be enabled for production.",
        "Plate recognition is a mock adapter. A live provider requires accuracy, bias, privacy, retention, and legal evaluation, and should retain human confirmation.",
        "Availability is estimated from software records and warden observations because the project uses no IoT occupancy sensors.",
        "Browser alerts depend on permission and application state; dependable background reminders require a production push-notification service.",
    ]:
        paragraph(doc, item, style="List Bullet")


def add_running_header_footer(doc: Document) -> None:
    for section in doc.sections:
        header = section.header
        header_p = header.paragraphs[0]
        header_p.text = ""
        header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = header_p.add_run("SmartPark Project Documentation")
        run.font.size = Pt(9)
        run.font.color.rgb = GRAY

        footer = section.footer
        footer_p = footer.paragraphs[0]
        footer_p.text = ""
        add_page_number(footer_p)
        for run in footer_p.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = GRAY


def scrub_metadata(path: Path) -> None:
    tmp = path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/core.xml":
                text = data.decode("utf-8")
                text = re.sub(r"<dc:creator>.*?</dc:creator>", "<dc:creator>SmartPark Project Team</dc:creator>", text)
                text = re.sub(
                    r"<cp:lastModifiedBy>.*?</cp:lastModifiedBy>",
                    "<cp:lastModifiedBy>SmartPark Project Team</cp:lastModifiedBy>",
                    text,
                )
                data = text.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)


def build() -> None:
    OUT_DIR.mkdir(exist_ok=True)
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_front_matter(doc)
    add_manual_toc(doc)

    for idx, chapter in enumerate(CHAPTERS):
        add_markdown_file(doc, chapter)
        if idx != len(CHAPTERS) - 1:
            doc.add_page_break()

    add_appendices(doc)
    add_running_header_footer(doc)
    doc.save(OUT_DOCX)
    scrub_metadata(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    build()
